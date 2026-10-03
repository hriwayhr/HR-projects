"""Workload.xlsx → документы базы для src/hr-metrics.html.

Листы «Юля», «Анна», «Катя»: блоки по месяцам, строка дат, строка дней недели,
затем строки «Количество …». Даты в шапке бывают datetime, числом 18.07,
строкой «18.07» или формулой «=D1+1» — берём только те, что попадают в месяц блока.

  python3 tools/import_workload.py Workload.xlsx > seed.json
"""
import datetime as dt, json, re, sys
import openpyxl

MONTHS = 'январь февраль март апрель май июнь июль август сентябрь октябрь ноябрь декабрь'.split()
# заголовок строки (нижний регистр, начало) → ключ метрики
METRICS = [
    ('количество открытых вакансий на начало', 'vac_start'),
    ('количество кандидатов на вакансию на начало', 'cand_start'),
    ('количество открытых вакансий на конец', 'vac_end'),
    ('количество кандидатов на вакансию на конец', 'cand_end'),
    ('количество собеседований (назнач', 'int_planned'),
    ('количество собеседований (факт', 'int_done'),
    ('количество звонков', 'calls'),
    ('количество первичных чатов', 'chats_first'),
    ('количество чатов первичн', 'chats_first'),
    ('количество чатов', 'chats'),
    ('количество новых', 'vac_published'),
    ('количество выходов', 'hires'),
    ('количество увольнений', 'dismissals'),
    ('количество отказов', 'rejections'),
    ('колчество отказов', 'rejections'),
]


def metric_key(label):
    s = ' '.join(str(label).lower().split())
    return next((k for p, k in METRICS if s.startswith(p)), None)


def parse_sheet(ws):
    out = {}
    year, month, dates = 2024, None, {}
    for row in ws.iter_rows(values_only=True):
        a = row[0]
        if isinstance(a, (int, float)) and 2020 < a < 2100:
            year = int(a)
            continue
        name = str(a).strip().lower() if a is not None else ''
        if name in MONTHS or name == '(':  # «(» — опечатка вместо «Июль» у Кати
            month = MONTHS.index(name) + 1 if name in MONTHS else 7
            dates = {}
            for i, v in enumerate(row[1:], 1):
                day = None
                if isinstance(v, dt.datetime):
                    day = v.day if v.month == month else None
                elif isinstance(v, (int, float, str)) and re.fullmatch(r'\d{1,2}\.\d{1,2}', str(v)):
                    d, m = str(v).split('.')
                    m = int(m.ljust(2, '0')) if len(m) == 1 else int(m)  # 18.1 → октябрь
                    day = int(d) if m == month else None
                elif isinstance(v, str) and v.startswith('=') and i - 1 in dates:
                    day = dates[i - 1] + 1
                if day:
                    dates[i] = day
            continue
        key = metric_key(a) if a else None
        if not key or not month:
            continue
        for i, day in dates.items():
            v = row[i] if i < len(row) else None
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                try:
                    iso = dt.date(year, month, day).isoformat()
                except ValueError:
                    continue
                out.setdefault(iso, {})[key] = v
    return out


PEOPLE = {'Юля': 'yulia', 'Анна': 'anna', 'Катя': 'katya'}


def to_docs(data):
    """{имя: {дата: метрики}} → документы базы hr-metrics.html: config/team и months/<pid>_<YYYY-MM>."""
    docs = {'config/team': {'people': {pid: {'name': n, 'order': i}
                                        for i, (n, pid) in enumerate(PEOPLE.items(), 1)}}}
    for name, days in data.items():
        pid = PEOPLE[name]
        for day, m in days.items():
            doc = docs.setdefault(f'months/{pid}_{day[:7]}', {'person': pid, 'month': day[:7], 'days': {}})
            doc['days'][day] = {k: int(v) for k, v in m.items()}
    return docs


VAC_COLS = ['title', 'dept', 'status', 'link', 'rate', 'requestedAt', 'publishedAt', 'archivedAt', 'manager',
            'extensions', 'closedAt', 'offersSent', 'offersAccepted', 'startAt', 'comment']
VAC_DATES = {'requestedAt', 'publishedAt', 'archivedAt', 'closedAt', 'startAt'}
VAC_NUMS = {'rate', 'extensions', 'offersSent', 'offersAccepted'}


def to_date(v):
    """datetime, «28.05.2025», «28.05.2025 вышел сотрудник» или серийный номер Excel → ISO; иначе None."""
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, (int, float)) and 30000 < v < 80000:
        return (dt.date(1899, 12, 30) + dt.timedelta(days=int(v))).isoformat()
    m = re.match(r'\s*(\d{1,2})\.(\d{1,2})\.(\d{4})', str(v or ''))
    return dt.date(int(m[3]), int(m[2]), int(m[1])).isoformat() if m else None


def parse_vacancies(ws):
    """Лист «вакансии в работе» → документы vacancies/v001…; текст из ячеек дат уходит в комментарий."""
    docs = {}
    for n, row in enumerate(ws.iter_rows(min_row=2, max_col=15, values_only=True), 1):
        if not row[0]:
            continue
        v, notes = {}, []
        for key, cell in zip(VAC_COLS, row):
            if cell is None or (isinstance(cell, str) and (not cell.strip() or cell.startswith('='))):
                continue
            if key in VAC_DATES:
                d = to_date(cell)
                if d:
                    v[key] = d
                if isinstance(cell, str) and not re.fullmatch(r'\s*\d{1,2}\.\d{1,2}\.\d{4}\s*', cell):
                    notes.append(cell.strip())
            elif key in VAC_NUMS:
                try:
                    v[key] = float(str(cell).replace(',', '.'))
                except ValueError:
                    notes.append(str(cell).strip())
            else:
                v[key] = ' '.join(str(cell).split())
        if notes:
            v['comment'] = '; '.join(([v['comment']] if 'comment' in v else []) + notes)
        docs[f'vacancies/v{n:03d}'] = v
    return docs


def main(path):
    wb = openpyxl.load_workbook(path)
    docs = to_docs({p: parse_sheet(wb[p]) for p in PEOPLE})
    docs.update(parse_vacancies(wb['вакансии в работе']))
    json.dump(docs, sys.stdout, ensure_ascii=False)


if __name__ == '__main__':
    main(sys.argv[1])
